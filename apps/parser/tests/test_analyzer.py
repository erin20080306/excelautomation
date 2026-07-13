from datetime import datetime
from pathlib import Path

import openpyxl

from app.analyzer import analyze_file, detect_header, parse_date, parse_number


def write_workbook(path: Path) -> None:
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = "本月銷售"
    sheet.merge_cells("A1:D1")
    sheet["A1"] = "2026 年銷售報表"
    sheet.append([None, None, None, None])
    sheet.append(["客戶名稱", "商品名稱", "QTY", "Total"])
    sheet.append(["甲公司", "A 商品", 2, 200])
    sheet.append(["乙公司", "B 商品", 3, 450])
    sheet.append(["總計", None, 5, 650])
    sheet.row_dimensions[2].hidden = True
    sheet["D4"].number_format = "$#,##0"
    workbook.save(path)


def test_detects_non_first_header_and_excludes_total(tmp_path: Path):
    target = tmp_path / "銷售報表.xlsx"
    write_workbook(target)
    result = analyze_file(str(target), target.name)
    sheet = result["workbook"]["sheets"][0]
    assert sheet["headerRow"] == 3
    assert sheet["dataRowCount"] == 2
    assert sheet["excludedRows"] == [{"row": 6, "type": "total"}]
    assert sheet["mergedCells"] == ["A1:D1"]
    assert sheet["hiddenRows"] == [2]
    assert result["classification"]["type"] == "sales"


def test_maps_aliases_and_preserves_typed_values(tmp_path: Path):
    target = tmp_path / "sales.xlsx"
    write_workbook(target)
    result = analyze_file(str(target), target.name)
    fields = result["workbook"]["sheets"][0]["fields"]
    assert [field["targetKey"] for field in fields] == ["customer", "product", "quantity", "amount"]
    row = result["workbook"]["sheets"][0]["normalizedRows"][0]
    assert row["quantity"] == 2
    assert row["amount"] == 200


def test_multilevel_header_detection():
    rows = [["銷售資料", "銷售資料", "客戶資料"], ["日期", "金額", "公司名稱"], ["2026/01/01", 100, "甲"]]
    index, levels, confidence = detect_header(rows)
    assert index in (0, 1)
    assert levels in (1, 2)
    assert confidence > 0.5


def test_date_and_money_normalization():
    assert parse_date("115/07/11") == datetime(2026, 7, 11)
    assert parse_date(2) == datetime(1900, 1, 1)
    assert parse_number("NTD 1,234.50") == 1234.5
    assert parse_number("(1,000)") == -1000
    assert parse_number("25%") == 0.25


def test_detects_deep_header_and_semantic_synonyms_without_template(tmp_path: Path):
    target = tmp_path / "不固定格式.xlsx"
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    for index in range(55):
        sheet.append([f"前言說明 {index + 1}", None, None, None])
    sheet.append(["訂單日期", "客戶簡稱", "品號", "銷售數量"])
    sheet.append(["2026/07/01", "甲公司", "A-01", 3])
    sheet.append(["2026/07/02", "乙公司", "B-02", 5])
    workbook.save(target)

    result = analyze_file(str(target), target.name)
    analyzed = result["workbook"]["sheets"][0]
    assert analyzed["headerRow"] == 56
    assert [field["targetKey"] for field in analyzed["fields"]] == ["date", "customer", "product_id", "quantity"]
    assert analyzed["dataRowCount"] == 2
