from pathlib import Path
from datetime import datetime

import openpyxl

from app.exporter import create_export
from app.analyzer import analyze_file


def test_export_is_openable_typed_and_formula_safe(tmp_path: Path):
    output = tmp_path / "result.xlsx"
    payload = {
        "config": {"includeProfessionalReport": True, "includeOverview": True, "includeMappings": True, "includeExceptions": True, "mergeAll": True, "mergeByType": False, "separateSourceSheets": True, "includeStatistics": True, "preserveRaw": True, "includeAudit": True},
        "analyses": [{
            "fileName": "sales.xlsx", "requiresReview": False, "warnings": [],
            "classification": {"type": "sales", "confidence": 0.9},
            "workbook": {"sheetCount": 1, "sheets": [{
                "name": "Sheet1", "dataRowCount": 1, "fields": [{"sourceName": "金額", "targetKey": "amount", "dataType": "number", "confidence": 0.98, "evidence": ["alias"]}],
                "normalizedRows": [{"date": "2026-07-11", "amount": 1200, "comment": "=HYPERLINK(\"bad\")"}],
                "rawRows": [{"日期": "2026/07/11", "金額": "1,200"}], "excludedRows": [], "changes": []
            }]}
        }]
    }
    create_export(payload, str(output))
    workbook = openpyxl.load_workbook(output, data_only=False)
    assert "01_檔案總覽" in workbook.sheetnames
    assert "00_專業分析" in workbook.sheetnames
    assert "整合總表" in workbook.sheetnames
    assert "統計報表" in workbook.sheetnames
    assert any(name.startswith("原始_") for name in workbook.sheetnames)
    sheet = workbook["整合總表"]
    headers = [cell.value for cell in sheet[1]]
    date_cell = sheet.cell(2, headers.index("date") + 1)
    amount_cell = sheet.cell(2, headers.index("amount") + 1)
    comment_cell = sheet.cell(2, headers.index("comment") + 1)
    assert isinstance(date_cell.value, datetime)
    assert amount_cell.value == 1200
    assert comment_cell.value.startswith("'=")
    assert workbook["00_專業分析"]["A1"].value == "ExcelMaster 專業整合分析報告"
    assert analyze_file(str(output), output.name)["workbook"]["sheetCount"] == len(workbook.sheetnames)
