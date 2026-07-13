from __future__ import annotations

import csv
import hashlib
import io
import json
import math
import re
import unicodedata
from collections import Counter
from datetime import date, datetime, timedelta
from pathlib import Path
from typing import Any, Iterable

import openpyxl
import pandas as pd
from openpyxl.cell.cell import Cell
from openpyxl.worksheet.worksheet import Worksheet

REPORT_KEYWORDS: dict[str, set[str]] = {
    "quotation": {"報價", "quotation", "quote", "有效期限"},
    "order": {"訂單", "order", "訂購", "訂單號"},
    "purchase": {"採購", "purchase", "供應商", "vendor", "po"},
    "sales": {"銷售", "sales", "出貨", "營業額", "客戶", "單價", "金額"},
    "inventory": {"庫存", "inventory", "倉庫", "安全庫存", "現有數量", "stock"},
    "attendance": {"出勤", "考勤", "上班", "下班", "工時", "attendance"},
    "payroll": {"薪資", "薪水", "salary", "payroll", "應發", "實發"},
    "customer": {"客戶名單", "customer", "聯絡人", "公司名稱", "電話", "地址"},
    "employee": {"員工名單", "employee", "員工編號", "部門", "職稱"},
    "product": {"商品清單", "產品", "product", "料號", "sku", "規格"},
    "finance": {"財務", "finance", "會計", "借方", "貸方", "科目", "損益"},
    "recruitment": {"招募", "應徵", "面試", "履歷", "recruitment"},
    "logistics": {"物流", "運單", "收件", "配送", "tracking", "logistics"},
    "reconciliation": {"對帳", "reconciliation", "帳款", "差異", "發票"},
}

FIELD_ALIASES: dict[str, set[str]] = {
    "date": {"日期", "交易日期", "訂單日期", "出貨日期", "採購日期", "建立日期", "時間", "date", "datetime", "createdat"},
    "customer": {"客戶", "客戶名稱", "公司", "公司名稱", "買方", "客戶簡稱", "customer", "client", "company"},
    "company": {"公司名稱", "客戶名稱", "company"},
    "contact": {"聯絡人", "contact", "窗口"},
    "product": {"商品", "商品名稱", "產品", "品項", "品名", "product", "item"},
    "product_id": {"商品編號", "產品編號", "品號", "貨號", "料號", "sku", "itemno", "itemcode", "productid"},
    "quantity": {"數量", "訂購數量", "銷售數量", "出貨量", "庫存量", "現有數量", "件數", "qty", "quantity", "onhand"},
    "unit_price": {"單價", "售價", "價格", "未稅單價", "含稅單價", "price", "unitprice"},
    "amount": {"金額", "總額", "未稅金額", "含稅金額", "銷售額", "小計", "合計", "total", "amount", "採購金額"},
    "name": {"姓名", "員工姓名", "人員姓名", "name", "employeename"},
    "employee_id": {"員工編號", "工號", "employeeid", "staffid"},
    "supplier": {"供應商", "廠商", "vendor", "supplier"},
    "phone": {"電話", "手機", "聯絡電話", "phone", "tel", "mobile"},
    "email": {"email", "電子郵件", "mail", "e-mail"},
    "address": {"地址", "公司地址", "address"},
    "warehouse": {"倉庫", "倉別", "庫別", "儲位", "warehouse", "location"},
    "unit": {"單位", "uom", "unit"},
    "specification": {"規格", "型號", "spec", "specification"},
    "purchase_no": {"採購單號", "po", "pono", "p/ono"},
    "order_no": {"訂單編號", "訂單號", "單號", "orderno", "orderid"},
    "invoice_no": {"發票號碼", "發票號", "invoice", "invoiceno"},
    "salesperson": {"業務", "業務員", "負責人", "承辦人", "salesperson", "owner"},
    "department": {"部門", "單位部門", "department", "dept"},
    "status": {"狀態", "處理狀態", "訂單狀態", "status"},
    "category": {"分類", "類別", "產品分類", "category", "type"},
    "tax": {"稅額", "營業稅", "tax", "vat"},
    "discount": {"折扣", "折讓", "discount"},
}

TOTAL_WORDS = {"合計", "總計", "總額", "grandtotal", "total"}
SUBTOTAL_WORDS = {"小計", "subtotal"}
NOTE_WORDS = {"備註", "說明", "註", "簽名", "核准", "製表", "note", "remark"}


def normalize_text(value: Any) -> str:
    if value is None:
        return ""
    text = unicodedata.normalize("NFKC", str(value)).strip().lower()
    return re.sub(r"[^\w\u3400-\u9fff]", "", text)


def text_similarity(left: str, right: str) -> float:
    if not left or not right:
        return 0.0
    if left == right:
        return 1.0
    if min(len(left), len(right)) >= 2 and (left in right or right in left):
        return 0.9 * min(len(left), len(right)) / max(len(left), len(right)) + 0.1
    def grams(value: str) -> set[str]:
        return {value[index:index + 2] for index in range(max(1, len(value) - 1))}
    left_grams, right_grams = grams(left), grams(right)
    return len(left_grams & right_grams) / max(1, len(left_grams | right_grams))


def serializable(value: Any) -> Any:
    if isinstance(value, (datetime, date)):
        return value.isoformat()
    if isinstance(value, float) and (math.isnan(value) or math.isinf(value)):
        return None
    if value is None or isinstance(value, (str, int, float, bool)):
        return value
    return str(value)


def classify_value(values: list[Any]) -> tuple[str, float]:
    present = [value for value in values if value not in (None, "")][:200]
    if not present:
        return "empty", 0.0
    counters: Counter[str] = Counter()
    for value in present:
        if isinstance(value, bool):
            counters["boolean"] += 1
        elif isinstance(value, (datetime, date)):
            counters["date"] += 1
        elif isinstance(value, (int, float)) and not isinstance(value, bool):
            counters["number"] += 1
        else:
            text = str(value).strip()
            if re.fullmatch(r"[^@\s]+@[^@\s]+\.[^@\s]+", text):
                counters["email"] += 1
            elif re.fullmatch(r"\+?[\d() -]{8,}", text):
                counters["phone"] += 1
            elif parse_date(text) is not None:
                counters["date"] += 1
            elif parse_number(text) is not None:
                counters["number"] += 1
            else:
                counters["text"] += 1
    value_type, count = counters.most_common(1)[0]
    return value_type, round(count / len(present), 3)


def parse_number(value: Any) -> float | int | None:
    if isinstance(value, bool):
        return None
    if isinstance(value, (int, float)):
        return value
    if not isinstance(value, str):
        return None
    text = unicodedata.normalize("NFKC", value).strip()
    if not text:
        return None
    negative = text.startswith("(") and text.endswith(")")
    text = re.sub(r"[,$€£¥NTD\s]", "", text).strip("()")
    percent = text.endswith("%")
    text = text.rstrip("%")
    try:
        number = float(text)
        number = -number if negative else number
        number = number / 100 if percent else number
        return int(number) if number.is_integer() and not percent else number
    except ValueError:
        return None


def parse_date(value: Any) -> datetime | None:
    if isinstance(value, datetime):
        return value
    if isinstance(value, date):
        return datetime.combine(value, datetime.min.time())
    if isinstance(value, (int, float)) and 1 <= value <= 2958465:
        try:
            return datetime(1899, 12, 30) + timedelta(days=float(value))
        except (OverflowError, ValueError):
            return None
    if not isinstance(value, str):
        return None
    text = unicodedata.normalize("NFKC", value).strip()
    roc = re.fullmatch(r"(\d{2,3})[./-](\d{1,2})[./-](\d{1,2})", text)
    if roc and int(roc.group(1)) < 1911:
        try:
            return datetime(int(roc.group(1)) + 1911, int(roc.group(2)), int(roc.group(3)))
        except ValueError:
            return None
    for fmt in ("%Y/%m/%d", "%Y-%m-%d", "%Y.%m.%d", "%d/%m/%Y", "%m/%d/%Y", "%Y%m%d"):
        try:
            return datetime.strptime(text, fmt)
        except ValueError:
            continue
    return None


def score_header(rows: list[list[Any]], index: int) -> float:
    row = rows[index]
    nonempty = [value for value in row if value not in (None, "")]
    if len(nonempty) < 2:
        return 0.0
    density = len(nonempty) / max(1, len(row))
    text_ratio = sum(isinstance(value, str) for value in nonempty) / len(nonempty)
    unique_ratio = len({normalize_text(value) for value in nonempty}) / len(nonempty)
    keyword_hits = sum(any(normalize_text(value) in {normalize_text(alias) for alias in aliases} for aliases in FIELD_ALIASES.values()) for value in nonempty)
    next_density = 0.0
    type_transition = 0.0
    if index + 1 < len(rows):
        next_values = [value for value in rows[index + 1] if value not in (None, "")]
        next_nonempty = len(next_values)
        next_density = next_nonempty / max(1, len(row))
        if next_values:
            header_text_ratio = sum(isinstance(value, str) for value in nonempty) / len(nonempty)
            next_text_ratio = sum(isinstance(value, str) for value in next_values) / len(next_values)
            type_transition = max(0.0, header_text_ratio - next_text_ratio)
    title_penalty = 0.35 if len(nonempty) == 1 else 0
    return max(0.0, min(1.0, 0.22 * density + 0.22 * text_ratio + 0.14 * unique_ratio + 0.22 * min(1, keyword_hits / 3) + 0.12 * next_density + 0.08 * type_transition - title_penalty))


def detect_header(rows: list[list[Any]]) -> tuple[int, int, float]:
    if not rows:
        return 0, 1, 0.0
    scores = [(index, score_header(rows, index)) for index in range(min(100, len(rows)))]
    header_index, confidence = max(scores, key=lambda item: item[1])
    levels = 1
    if header_index > 0 and score_header(rows, header_index - 1) >= confidence * 0.82:
        above = rows[header_index - 1]
        if sum(value not in (None, "") for value in above) >= 2:
            header_index -= 1
            levels = 2
            confidence = min(1.0, confidence + 0.05)
    return header_index, levels, round(confidence, 3)


def flatten_headers(rows: list[list[Any]], header_index: int, levels: int, width: int) -> list[str]:
    headers: list[str] = []
    inherited: list[Any] = [None] * width
    for col in range(width):
        parts: list[str] = []
        for level in range(levels):
            value = rows[header_index + level][col] if col < len(rows[header_index + level]) else None
            if value in (None, "") and level == 0 and col > 0:
                value = inherited[col - 1]
            if value not in (None, ""):
                inherited[col] = value
                text = str(value).strip()
                if not parts or parts[-1] != text:
                    parts.append(text)
        headers.append(" / ".join(parts) if parts else f"未命名欄位_{col + 1}")
    seen: Counter[str] = Counter()
    unique: list[str] = []
    for header in headers:
        seen[header] += 1
        unique.append(header if seen[header] == 1 else f"{header}_{seen[header]}")
    return unique


def row_kind(row: list[Any]) -> str:
    normalized = {normalize_text(value) for value in row if value not in (None, "")}
    if normalized & TOTAL_WORDS:
        return "total"
    if normalized & SUBTOTAL_WORDS:
        return "subtotal"
    if normalized & NOTE_WORDS and len(normalized) <= 5:
        return "note"
    return "data"


def map_field(header: str, samples: list[Any], column_index: int, total_columns: int) -> dict[str, Any]:
    normalized = normalize_text(header)
    exact = [(key, aliases) for key, aliases in FIELD_ALIASES.items() if normalized in aliases]
    value_type, type_confidence = classify_value(samples)
    if exact:
        target = exact[0][0]
        confidence = 0.98
        evidence = ["欄位別名完全符合"]
    else:
        candidates: list[tuple[str, float]] = []
        for key, aliases in FIELD_ALIASES.items():
            similarity = max((text_similarity(normalized, normalize_text(alias)) for alias in aliases), default=0)
            if similarity >= 0.58:
                candidates.append((key, min(0.92, similarity * 0.9)))
        if value_type in {"email", "phone", "date"}:
            typed_target = {"email": "email", "phone": "phone", "date": "date"}[value_type]
            candidates.append((typed_target, 0.72 * type_confidence))
        if candidates:
            target, confidence = max(candidates, key=lambda item: item[1])
            evidence = ["正規化文字相似度或資料型態推論"]
        else:
            target, confidence, evidence = None, 0.25, ["沒有可靠的規則對應"]
    if target == "quantity" and column_index + 1 < total_columns:
        confidence = min(1.0, confidence + 0.02)
        evidence.append("欄位位置關聯")
    return {
        "sourceName": header, "normalizedName": normalized, "columnIndex": column_index + 1,
        "targetKey": target, "dataType": value_type, "confidence": round(confidence, 3),
        "evidence": evidence, "requiresReview": confidence < 0.75,
    }


def clean_value(value: Any, field: dict[str, Any]) -> tuple[Any, str | None]:
    if value is None:
        return None, None
    original = serializable(value)
    target = field.get("targetKey")
    data_type = field.get("dataType")
    if isinstance(value, str):
        value = unicodedata.normalize("NFKC", value).replace("\u200b", "").strip()
    rule = None
    if target == "date" or data_type == "date":
        parsed = parse_date(value)
        if parsed:
            value, rule = parsed.isoformat(), "date_standardization"
    elif target in {"quantity", "unit_price", "amount"} or data_type == "number":
        parsed_number = parse_number(value)
        if parsed_number is not None:
            value, rule = parsed_number, "number_standardization"
    elif target == "phone" and isinstance(value, str):
        value, rule = re.sub(r"[^\d+]", "", value), "phone_standardization"
    if serializable(value) == original:
        rule = None
    return serializable(value), rule


def worksheet_rows(ws: Worksheet) -> list[list[Any]]:
    return [[serializable(cell.value) for cell in row] for row in ws.iter_rows(min_row=1, max_row=ws.max_row, max_col=ws.max_column)]


def cell_format(cell: Cell) -> dict[str, Any]:
    border_styles = {}
    for side in ("left", "right", "top", "bottom"):
        border_side = getattr(cell.border, side, None)
        border_styles[side] = border_side.style if border_side is not None else None
    return {
        "coordinate": cell.coordinate,
        "numberFormat": cell.number_format,
        "font": {"name": cell.font.name, "size": cell.font.sz, "bold": bool(cell.font.bold), "italic": bool(cell.font.italic), "color": cell.font.color.rgb if cell.font.color and cell.font.color.type == "rgb" else None},
        "fill": cell.fill.fgColor.rgb if cell.fill and cell.fill.fgColor.type == "rgb" else None,
        "border": border_styles,
        "formula": cell.value if cell.data_type == "f" else None,
    }


def analyze_sheet(ws: Worksheet) -> dict[str, Any]:
    rows = worksheet_rows(ws)
    width = max((len(row) for row in rows), default=0)
    header_index, header_levels, header_confidence = detect_header(rows)
    headers = flatten_headers(rows, header_index, header_levels, width) if width else []
    data_start = header_index + header_levels
    last_nonempty = max((index for index, row in enumerate(rows) if any(value not in (None, "") for value in row)), default=data_start)
    kinds = {index: row_kind(rows[index]) for index in range(data_start, last_nonempty + 1)}
    normalized_headers = {normalize_text(header.split(" / ")[-1]) for header in headers}
    for index in range(data_start, last_nonempty + 1):
        values = [normalize_text(value) for value in rows[index] if value not in (None, "")]
        if values and len(set(values) & normalized_headers) / len(values) >= 0.7:
            kinds[index] = "repeated_header"
        elif width >= 3 and len(values) == 1 and not any(parse_number(value) is not None or parse_date(value) is not None for value in rows[index] if value not in (None, "")):
            kinds[index] = "note"
    data_indices = [index for index, kind in kinds.items() if kind == "data" and any(value not in (None, "") for value in rows[index])]
    data_end = max(data_indices, default=data_start - 1)
    fields = [map_field(header, [rows[index][column] if column < len(rows[index]) else None for index in data_indices], column, width) for column, header in enumerate(headers)]
    normalized_rows: list[dict[str, Any]] = []
    raw_rows: list[dict[str, Any]] = []
    changes: list[dict[str, Any]] = []
    for index in data_indices:
        output: dict[str, Any] = {"_sourceRow": index + 1}
        raw_output: dict[str, Any] = {"_sourceRow": index + 1}
        for column, field in enumerate(fields):
            raw = rows[index][column] if column < len(rows[index]) else None
            value, rule = clean_value(raw, field)
            key = field.get("targetKey") or field["sourceName"]
            if key in output:
                key = f"{key}_{column + 1}"
            output[key] = value
            raw_output[field["sourceName"]] = serializable(raw)
            if rule:
                changes.append({"cell": f"{openpyxl.utils.get_column_letter(column + 1)}{index + 1}", "raw": serializable(raw), "standard": value, "rule": rule, "humanModified": False})
        normalized_rows.append(output)
        raw_rows.append(raw_output)
    styled_cells = [cell_format(cell) for row in ws.iter_rows() for cell in row if cell.value is not None and (cell.has_style or cell.data_type == "f")]
    return {
        "name": ws.title,
        "position": ws.parent.index(ws),
        "hidden": ws.sheet_state != "visible",
        "maxRow": ws.max_row,
        "maxColumn": ws.max_column,
        "usedRange": f"A1:{openpyxl.utils.get_column_letter(max(1, ws.max_column))}{max(1, ws.max_row)}",
        "headerRow": header_index + 1,
        "headerLevels": header_levels,
        "headerConfidence": header_confidence,
        "dataStartRow": data_start + 1,
        "dataEndRow": data_end + 1,
        "dataRowCount": len(normalized_rows),
        "mergedCells": [str(item) for item in ws.merged_cells.ranges],
        "hiddenRows": [index for index, dim in ws.row_dimensions.items() if dim.hidden],
        "hiddenColumns": [index for index, dim in ws.column_dimensions.items() if dim.hidden],
        "formulaCount": sum(1 for row in ws.iter_rows() for cell in row if cell.data_type == "f"),
        "regions": [{"type": "data", "startRow": data_start + 1, "endRow": data_end + 1, "startColumn": 1, "endColumn": width, "confidence": header_confidence}],
        "excludedRows": [{"row": index + 1, "type": kind} for index, kind in kinds.items() if kind != "data"],
        "fields": fields,
        "normalizedRows": normalized_rows,
        "rawRows": raw_rows,
        "changes": changes,
        "styles": styled_cells,
    }


def analyze_tabular_sheets(sheets: dict[str, list[list[Any]]]) -> list[dict[str, Any]]:
    workbook = openpyxl.Workbook()
    workbook.remove(workbook.active)
    for name, rows in sheets.items():
        ws = workbook.create_sheet(str(name)[:31] or "Sheet")
        for row in rows:
            ws.append([serializable(value) for value in row])
    return [analyze_sheet(sheet) for sheet in workbook.worksheets]


def classify(file_name: str, sheets: list[dict[str, Any]]) -> dict[str, Any]:
    haystack = " ".join([file_name, *[sheet["name"] for sheet in sheets], *[field["sourceName"] for sheet in sheets for field in sheet["fields"]]]).lower()
    normalized = normalize_text(haystack)
    scores: dict[str, float] = {}
    reasons: dict[str, list[str]] = {}
    for report_type, keywords in REPORT_KEYWORDS.items():
        hits = [keyword for keyword in keywords if normalize_text(keyword) in normalized]
        score = min(0.98, len(hits) * 0.16 + (0.08 if any(normalize_text(keyword) in normalize_text(file_name) for keyword in keywords) else 0))
        scores[report_type] = score
        reasons[report_type] = [f"符合關鍵字：{', '.join(hits[:6])}"] if hits else []
    ordered = sorted(scores.items(), key=lambda item: item[1], reverse=True)
    winner, confidence = ordered[0] if ordered else ("unknown", 0.0)
    if confidence < 0.35:
        winner = "unknown"
        reasons[winner] = ["規則證據不足，需要人工確認"]
        confidence = max(0.2, confidence)
    return {"type": winner, "confidence": round(confidence, 3), "reasons": reasons.get(winner, []), "candidates": [{"type": key, "confidence": round(score, 3)} for key, score in ordered[:3]]}


def analyze_file(file_path: str, original_name: str) -> dict[str, Any]:
    extension = Path(original_name).suffix.lower()
    file_hash = hashlib.sha256(Path(file_path).read_bytes()).hexdigest()
    has_macros = extension == ".xlsm"
    if extension in {".xlsx", ".xlsm"}:
        workbook = openpyxl.load_workbook(file_path, data_only=False, read_only=False, keep_vba=has_macros, keep_links=False)
        sheets = [analyze_sheet(sheet) for sheet in workbook.worksheets]
    elif extension == ".xls":
        frames = pd.read_excel(file_path, sheet_name=None, header=None, dtype=object)
        sheets = analyze_tabular_sheets({name: frame.where(pd.notnull(frame), None).values.tolist() for name, frame in frames.items()})
    elif extension in {".csv", ".tsv"}:
        raw = Path(file_path).read_bytes()
        text = None
        for encoding in ("utf-8-sig", "big5", "cp950"):
            try:
                text = raw.decode(encoding)
                break
            except UnicodeDecodeError:
                continue
        if text is None:
            raise ValueError("CSV 編碼無法辨識")
        delimiter = "\t" if extension == ".tsv" else ","
        rows = list(csv.reader(io.StringIO(text), delimiter=delimiter))
        sheets = analyze_tabular_sheets({Path(original_name).stem[:31]: rows})
    else:
        raise ValueError("不支援的檔案格式")
    classification = classify(original_name, sheets)
    all_fields = [field for sheet in sheets for field in sheet["fields"]]
    warnings: list[str] = []
    if has_macros:
        warnings.append("檔案含 VBA；系統不會執行巨集，只讀取工作表資料")
    if any(sheet["headerConfidence"] < 0.65 for sheet in sheets):
        warnings.append("部分工作表的表頭信心度偏低")
    if any(field["requiresReview"] for field in all_fields):
        warnings.append("部分欄位需要人工確認")
    requires_review = classification["confidence"] < 0.7 or bool(warnings)
    return {
        "fileName": original_name,
        "fileHash": file_hash,
        "workbook": {"sheetCount": len(sheets), "hasMacros": has_macros, "sheets": sheets},
        "classification": classification,
        "fields": all_fields,
        "warnings": warnings,
        "requiresReview": requires_review,
    }
