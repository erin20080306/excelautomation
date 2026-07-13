from __future__ import annotations

import re
from datetime import datetime
from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter


INVALID_SHEET_CHARS = re.compile(r"[\\/*?:\[\]]")
FORMULA_PREFIXES = ("=", "+", "-", "@", "\t", "\r")
REPORT_TYPE_LABELS = {
    "quotation": "報價單", "order": "訂單", "purchase": "採購", "sales": "銷售",
    "inventory": "庫存", "attendance": "出勤", "payroll": "薪資", "unknown": "一般資料",
}


def safe_sheet_name(name: str, used: set[str]) -> str:
    base = INVALID_SHEET_CHARS.sub("_", name).strip("'")[:31] or "資料"
    candidate, number = base, 2
    while candidate in used:
        suffix = f"_{number}"
        candidate = f"{base[:31-len(suffix)]}{suffix}"
        number += 1
    used.add(candidate)
    return candidate


def safe_value(value: Any) -> Any:
    if isinstance(value, str):
        if value.startswith(FORMULA_PREFIXES):
            return "'" + value
        try:
            if re.fullmatch(r"\d{4}-\d{2}-\d{2}(T.*)?", value):
                return datetime.fromisoformat(value)
        except ValueError:
            pass
    return value


def write_table(ws, headers: list[str], rows: list[dict[str, Any]]) -> None:
    if not headers:
        headers = ["資料"]
    ws.append(headers)
    for row in rows:
        ws.append([safe_value(row.get(header)) for header in headers])
    header_fill = PatternFill("solid", fgColor="0F766E")
    header_font = Font(color="FFFFFF", bold=True)
    thin = Side(style="thin", color="D7E3E1")
    for cell in ws[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = Border(bottom=thin)
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    for column, header in enumerate(headers, 1):
        values = [str(row.get(header, "")) for row in rows[:200]]
        ws.column_dimensions[get_column_letter(column)].width = min(36, max(10, len(header) + 2, *(len(value) + 2 for value in values)))
        if ws.max_row >= 2:
            for cell in ws.iter_cols(min_col=column, max_col=column, min_row=2):
                item = cell[0]
                if isinstance(item.value, datetime):
                    item.number_format = "yyyy-mm-dd"
                elif isinstance(item.value, (int, float)):
                    item.number_format = "#,##0.00"


def create_professional_report(wb: Workbook, used: set[str], analyses: list[dict[str, Any]]) -> None:
    ws = wb.create_sheet(safe_sheet_name("00_專業分析", used))
    sheets = [(analysis, sheet) for analysis in analyses for sheet in analysis.get("workbook", {}).get("sheets", [])]
    fields = [field for _, sheet in sheets for field in sheet.get("fields", [])]
    total_rows = sum(sheet.get("dataRowCount", 0) for _, sheet in sheets)
    mapped_fields = sum(bool(field.get("targetKey")) for field in fields)
    mapping_rate = mapped_fields / len(fields) if fields else 0
    header_confidence = sum(sheet.get("headerConfidence", 0) for _, sheet in sheets) / len(sheets) if sheets else 0
    field_confidence = sum(field.get("confidence", 0) for field in fields) / len(fields) if fields else 0
    warning_count = sum(len(analysis.get("warnings", [])) for analysis in analyses)
    quality_score = max(0, min(1, header_confidence * 0.45 + field_confidence * 0.35 + mapping_rate * 0.2 - min(0.2, warning_count * 0.02)))

    ws.sheet_view.showGridLines = False
    ws.merge_cells("A1:H1")
    ws["A1"] = "ExcelMaster 專業整合分析報告"
    ws["A1"].font = Font(size=20, bold=True, color="FFFFFF")
    ws["A1"].fill = PatternFill("solid", fgColor="0F766E")
    ws["A1"].alignment = Alignment(vertical="center")
    ws.row_dimensions[1].height = 34
    ws.merge_cells("A2:H2")
    ws["A2"] = f"產生時間：{datetime.now():%Y-%m-%d %H:%M}　｜　系統依內容自動辨識表頭、欄位語意與資料型態，不依賴固定模板"
    ws["A2"].font = Font(color="475569", italic=True)

    kpis = [
        ("來源檔案", len(analyses), "份"), ("工作表", len(sheets), "張"),
        ("整合資料", total_rows, "筆"), ("品質分數", quality_score, ""),
    ]
    for index, (label, value, unit) in enumerate(kpis):
        column = 1 + index * 2
        ws.cell(4, column, label)
        ws.cell(4, column).font = Font(size=10, bold=True, color="64748B")
        ws.merge_cells(start_row=5, start_column=column, end_row=6, end_column=column + 1)
        cell = ws.cell(5, column)
        cell.value = value
        cell.font = Font(size=22, bold=True, color="0F172A")
        cell.fill = PatternFill("solid", fgColor="ECFDF5")
        cell.alignment = Alignment(horizontal="center", vertical="center")
        if label == "品質分數":
            cell.number_format = "0%"
        elif unit:
            cell.number_format = f'#,##0" {unit}"'

    summary_headers = ["來源檔案", "辨識類型", "工作表", "自動表頭", "資料筆數", "欄位對應率", "平均信心度", "狀態"]
    summary_rows: list[dict[str, Any]] = []
    for analysis in analyses:
        analysis_sheets = analysis.get("workbook", {}).get("sheets", [])
        analysis_fields = [field for sheet in analysis_sheets for field in sheet.get("fields", [])]
        mapped = sum(bool(field.get("targetKey")) for field in analysis_fields)
        confidence_values = [sheet.get("headerConfidence", 0) for sheet in analysis_sheets] + [field.get("confidence", 0) for field in analysis_fields]
        summary_rows.append({
            "來源檔案": analysis.get("fileName", ""),
            "辨識類型": REPORT_TYPE_LABELS.get(analysis.get("classification", {}).get("type"), analysis.get("classification", {}).get("type", "一般資料")),
            "工作表": len(analysis_sheets),
            "自動表頭": "、".join(f"{sheet.get('name')} 第 {sheet.get('headerRow', 1)} 列" for sheet in analysis_sheets),
            "資料筆數": sum(sheet.get("dataRowCount", 0) for sheet in analysis_sheets),
            "欄位對應率": mapped / len(analysis_fields) if analysis_fields else 0,
            "平均信心度": sum(confidence_values) / len(confidence_values) if confidence_values else 0,
            "狀態": "需要確認" if analysis.get("requiresReview") else "可直接使用",
        })
    start_row = 9
    for column, header in enumerate(summary_headers, 1):
        ws.cell(start_row, column, header)
    for row_index, row in enumerate(summary_rows, start_row + 1):
        for column, header in enumerate(summary_headers, 1):
            ws.cell(row_index, column, safe_value(row.get(header)))
    header_fill = PatternFill("solid", fgColor="115E59")
    for cell in ws[start_row]:
        cell.fill = header_fill
        cell.font = Font(color="FFFFFF", bold=True)
        cell.alignment = Alignment(horizontal="center")
    for row_index in range(start_row + 1, start_row + 1 + len(summary_rows)):
        ws.cell(row_index, 5).number_format = "#,##0"
        ws.cell(row_index, 6).number_format = "0%"
        ws.cell(row_index, 7).number_format = "0%"
        status_cell = ws.cell(row_index, 8)
        status_cell.fill = PatternFill("solid", fgColor="FEF3C7" if status_cell.value == "需要確認" else "DCFCE7")

    insight_row = start_row + len(summary_rows) + 3
    ws.merge_cells(start_row=insight_row, start_column=1, end_row=insight_row, end_column=8)
    ws.cell(insight_row, 1, "系統判讀與建議")
    ws.cell(insight_row, 1).font = Font(size=14, bold=True, color="0F172A")
    insights = [
        f"已自動辨識 {len(sheets)} 張工作表的表頭位置，平均信心度 {header_confidence:.0%}。",
        f"共辨識 {len(fields)} 個欄位，其中 {mapped_fields} 個已對應標準語意（{mapping_rate:.0%}）。",
        f"資料品質綜合分數為 {quality_score:.0%}；低信心度項目已列入「需要確認」而不會靜默覆蓋。",
    ]
    if warning_count:
        insights.append(f"偵測到 {warning_count} 項警示，請查看「03_異常資料」後再正式使用。")
    else:
        insights.append("未偵測到需特別處理的警示，可直接使用整合總表。")
    for offset, insight in enumerate(insights, 1):
        ws.merge_cells(start_row=insight_row + offset, start_column=1, end_row=insight_row + offset, end_column=8)
        ws.cell(insight_row + offset, 1, f"• {insight}")
        ws.cell(insight_row + offset, 1).alignment = Alignment(wrap_text=True, vertical="top")
        ws.cell(insight_row + offset, 1).font = Font(color="334155")

    if summary_rows:
        chart = BarChart()
        chart.title = "各來源有效資料筆數"
        chart.y_axis.title = "筆數"
        chart.height = 6.5
        chart.width = 12
        chart.add_data(Reference(ws, min_col=5, min_row=start_row, max_row=start_row + len(summary_rows)), titles_from_data=True)
        chart.set_categories(Reference(ws, min_col=1, min_row=start_row + 1, max_row=start_row + len(summary_rows)))
        ws.add_chart(chart, "J4")
    widths = {"A": 28, "B": 14, "C": 10, "D": 36, "E": 14, "F": 15, "G": 15, "H": 14}
    for column, width in widths.items():
        ws.column_dimensions[column].width = width
    ws.freeze_panes = "A10"


def create_export(payload: dict[str, Any], output_path: str) -> None:
    wb = Workbook()
    wb.remove(wb.active)
    used: set[str] = set()
    config = payload.get("config", {})
    analyses = payload.get("analyses", [])

    if config.get("includeProfessionalReport", True):
        create_professional_report(wb, used, analyses)

    if config.get("includeOverview", True):
        ws = wb.create_sheet(safe_sheet_name("01_檔案總覽", used))
        headers = ["來源檔案", "工作表數", "報表類型", "信心度", "資料筆數", "需確認"]
        rows = []
        for analysis in analyses:
            rows.append({
                "來源檔案": analysis["fileName"], "工作表數": analysis["workbook"]["sheetCount"],
                "報表類型": analysis["classification"]["type"], "信心度": analysis["classification"]["confidence"],
                "資料筆數": sum(sheet["dataRowCount"] for sheet in analysis["workbook"]["sheets"]),
                "需確認": "是" if analysis["requiresReview"] else "否",
            })
        write_table(ws, headers, rows)
        for cell in ws["D"][1:]:
            cell.number_format = "0.0%"

    if config.get("includeMappings", True):
        ws = wb.create_sheet(safe_sheet_name("02_欄位對應", used))
        headers = ["來源檔案", "工作表", "來源欄位", "標準欄位", "資料型態", "信心度", "判斷依據"]
        rows = []
        for analysis in analyses:
            for sheet in analysis["workbook"]["sheets"]:
                for field in sheet["fields"]:
                    rows.append({"來源檔案": analysis["fileName"], "工作表": sheet["name"], "來源欄位": field["sourceName"], "標準欄位": field.get("targetKey") or "未對應", "資料型態": field["dataType"], "信心度": field["confidence"], "判斷依據": "、".join(field["evidence"])})
        write_table(ws, headers, rows)
        for cell in ws["F"][1:]:
            cell.number_format = "0.0%"

    if config.get("includeExceptions", True):
        ws = wb.create_sheet(safe_sheet_name("03_異常資料", used))
        headers = ["來源檔案", "類型", "說明", "位置"]
        rows = []
        for analysis in analyses:
            rows.extend({"來源檔案": analysis["fileName"], "類型": "warning", "說明": warning, "位置": ""} for warning in analysis.get("warnings", []))
            for sheet in analysis["workbook"]["sheets"]:
                rows.extend({"來源檔案": analysis["fileName"], "類型": row["type"], "說明": "已排除非明細列", "位置": f"{sheet['name']}!{row['row']}"} for row in sheet.get("excludedRows", []))
        write_table(ws, headers, rows)

    grouped: dict[str, list[dict[str, Any]]] = {}
    output_sets: list[dict[str, Any]] = []
    data_sets = payload.get("dataSets") or []
    if data_sets:
        for data_set in data_sets:
            rows = data_set.get("rows", [])
            headers = list(dict.fromkeys(key for row in rows for key in row.keys()))
            write_table(wb.create_sheet(safe_sheet_name(data_set.get("name", "整合資料"), used)), headers, rows)
            output_sets.append({"name": data_set.get("name", "整合資料"), "rows": rows})
    elif config.get("mergeAll", True):
        rows = []
        for analysis in analyses:
            for sheet in analysis["workbook"]["sheets"]:
                rows.extend({"_來源檔案": analysis["fileName"], "_來源工作表": sheet["name"], **row} for row in sheet["normalizedRows"])
        headers = list(dict.fromkeys(key for row in rows for key in row.keys()))
        write_table(wb.create_sheet(safe_sheet_name("整合總表", used)), headers, rows)
        output_sets.append({"name": "整合總表", "rows": rows})
    elif config.get("mergeByType", True):
        for analysis in analyses:
            grouped.setdefault(analysis["classification"]["type"], []).extend({"_來源檔案": analysis["fileName"], "_來源工作表": sheet["name"], **row} for sheet in analysis["workbook"]["sheets"] for row in sheet["normalizedRows"])
        for report_type, rows in grouped.items():
            headers = list(dict.fromkeys(key for row in rows for key in row.keys()))
            write_table(wb.create_sheet(safe_sheet_name(f"{report_type}_總表", used)), headers, rows)
            output_sets.append({"name": f"{report_type}_總表", "rows": rows})

    if config.get("separateSourceSheets", True) or (not data_sets and not config.get("mergeByType", True)):
        for analysis in analyses:
            for sheet in analysis["workbook"]["sheets"]:
                rows = sheet["normalizedRows"]
                headers = list(dict.fromkeys(key for row in rows for key in row.keys()))
                write_table(wb.create_sheet(safe_sheet_name(f"{Path(analysis['fileName']).stem}_{sheet['name']}", used)), headers, rows)

    if config.get("preserveRaw", False):
        for analysis in analyses:
            for sheet in analysis["workbook"]["sheets"]:
                rows = sheet.get("rawRows", [])
                headers = list(dict.fromkeys(key for row in rows for key in row.keys()))
                write_table(wb.create_sheet(safe_sheet_name(f"原始_{Path(analysis['fileName']).stem}_{sheet['name']}", used)), headers, rows)

    split_by = config.get("splitBy")
    if split_by:
        split_sources = grouped.items() if grouped else ((data_set["name"], data_set["rows"]) for data_set in output_sets)
        for report_type, rows in split_sources:
            groups: dict[str, list[dict[str, Any]]] = {}
            for row in rows:
                groups.setdefault(str(row.get(split_by, "未分類")), []).append(row)
            for key, split_rows in groups.items():
                headers = list(dict.fromkeys(field for row in split_rows for field in row.keys()))
                write_table(wb.create_sheet(safe_sheet_name(f"{report_type}_{key}", used)), headers, split_rows)

    if config.get("includeStatistics", False):
        ws = wb.create_sheet(safe_sheet_name("統計報表", used))
        stats = []
        for data_set in output_sets:
            rows = data_set["rows"]
            fields = list(dict.fromkeys(key for row in rows for key in row.keys()))
            numeric_cells = sum(isinstance(value, (int, float)) and not isinstance(value, bool) for row in rows for value in row.values())
            stats.append({"資料集": data_set["name"], "資料筆數": len(rows), "欄位數": len(fields), "數值儲存格": numeric_cells})
        write_table(ws, ["資料集", "資料筆數", "欄位數", "數值儲存格"], stats)

    if config.get("includeAudit", True):
        ws = wb.create_sheet(safe_sheet_name("處理紀錄", used))
        headers = ["時間", "來源檔案", "清洗儲存格", "原始值", "標準值", "規則", "人工修改"]
        rows = []
        now = datetime.now()
        for analysis in analyses:
            for sheet in analysis["workbook"]["sheets"]:
                rows.extend({"時間": now, "來源檔案": analysis["fileName"], "清洗儲存格": change["cell"], "原始值": change["raw"], "標準值": change["standard"], "規則": change["rule"], "人工修改": "是" if change["humanModified"] else "否"} for change in sheet.get("changes", []))
        write_table(ws, headers, rows)

    if not wb.worksheets:
        wb.create_sheet("整合結果")
    wb.save(output_path)
